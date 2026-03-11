from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from app.llm.prompt_store import render_prompt
from app.review.claude_review import (
    _build_word_line_index,
    _enrich_report_evidence_locations,
    _has_precise_location_hint,
    run_bid_review_with_claude,
)


class _SeqClient:
    def __init__(self, outputs: list[str], tool_uses: list[dict[str, Any]] | None = None) -> None:
        self.timeout_sec = 120
        self._outputs = outputs
        self._tool_uses = tool_uses or [{"name": "Read", "input": {"path": "x"}}]
        self.prompts: list[str] = []
        self.ask_text_calls = 0

    def ask_text(self, prompt: str, *, task_label: str | None = None) -> str:
        idx = self.ask_text_calls
        self.ask_text_calls += 1
        self.prompts.append(prompt)
        return self._outputs[min(idx, len(self._outputs) - 1)]

    def ask_json(
        self,
        prompt: str,
        *,
        required_top_keys: list[str] | None = None,
        max_retries: int = 2,
        task_label: str | None = None,
    ) -> dict[str, Any] | list[Any]:
        raise AssertionError("ask_json should not be called in these tests")

    def get_last_tool_calls(self) -> list[str]:
        return [str(x.get("name", "")) for x in self._tool_uses]

    def get_last_tool_uses(self) -> list[dict[str, Any]]:
        return list(self._tool_uses)


def test_review_prompts_require_precise_page_or_section_lines() -> None:
    main_prompt = render_prompt(
        "review_main.md",
        workspace_dir="D:/code/docs",
        tender_stem="招标文件",
        bid_stem="投标文件",
        tender_path="D:/code/docs/招标文件.pdf",
        bid_path="D:/code/docs/投标文件.docx",
        user_instruction="无",
        instruction="无",
        tender_document_map="- 总页数: 100",
        bid_document_map="- 检测到的正文/模板块数量: 10",
        minimum_requirement_count="10",
    )
    second_prompt = render_prompt(
        "review_second_pass.md",
        workspace_dir="D:/code/docs",
        tender_stem="招标文件",
        bid_stem="投标文件",
        tender_path="D:/code/docs/招标文件.pdf",
        bid_path="D:/code/docs/投标文件.docx",
        user_instruction="无",
        initial_json="{}",
    )

    assert "招标文件第X页 Lm-Ln" in main_prompt
    assert "投标文件《第X章/第X节/章节名》Lm-Ln" in main_prompt
    assert "图片OCR Lm-Ln" in main_prompt
    assert "招标文件第X页 Lm-Ln" in second_prompt
    assert "投标文件《第X章/第X节/章节名》Lm-Ln" in second_prompt


def test_precise_location_hint_requires_line_level_precision() -> None:
    assert _has_precise_location_hint("招标文件第12页 L34-L39：报价精确到小数点后两位")
    assert _has_precise_location_hint("投标文件《资格审查申请书》L2-L3：申请人：示例公司")
    assert _has_precise_location_hint("投标文件第19页图片OCR L3-L7：开户银行")
    assert not _has_precise_location_hint("第12页：报价精确到小数点后两位")
    assert not _has_precise_location_hint("资格审查申请书P3显示申请人：示例公司")


def test_enrich_word_evidence_locations_adds_section_and_lines(tmp_path: Path) -> None:
    bid_path = tmp_path / "bid.docx"
    doc = Document()
    doc.add_paragraph("商务投标文件")
    doc.add_paragraph("资格审查申请书")
    doc.add_paragraph("申请人：测试科技有限公司")
    doc.add_paragraph("法定代表人：张三")
    doc.save(str(bid_path))

    report = {
        "findings": [
            {
                "id": "F001",
                "requirement_id": "R001",
                "status": "non_compliant",
                "issue": "申请人填写错误",
                "tender_evidence": "招标文件第1页 L1-L2：资格审查申请书应由投标人填写",
                "bid_evidence": "资格审查申请书显示申请人：测试科技有限公司",
                "recommendation": "更正申请人名称",
            }
        ]
    }

    enriched = _enrich_report_evidence_locations(
        report,
        tender_path=tmp_path / "tender.pdf",
        bid_path=bid_path,
    )

    evidence = enriched["findings"][0]["bid_evidence"]
    assert evidence.startswith("投标文件《资格审查申请书》L2-L2：")
    assert "申请人：测试科技有限公司" in evidence


def test_build_word_line_index_keeps_paragraph_and_table_order(tmp_path: Path) -> None:
    bid_path = tmp_path / "bid.docx"
    doc = Document()
    doc.add_paragraph("第一章 总则")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "表格字段"
    table.rows[0].cells[1].text = "测试值"
    doc.add_paragraph("收尾段落")
    doc.save(str(bid_path))

    index = _build_word_line_index(str(bid_path))

    assert [item["text"] for item in index[:3]] == [
        "第一章 总则",
        "表格字段 | 测试值",
        "收尾段落",
    ]
    assert [item["line_no"] for item in index[:3]] == [1, 2, 3]


def test_run_bid_review_retries_when_precise_locations_still_missing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv("BID_REVIEW_ENABLE_SECOND_PASS", "0")
    monkeypatch.setenv("BID_REVIEW_DOCX_OCR_REQUIRED", "0")
    monkeypatch.setenv("BID_REVIEW_ENFORCE_COMPLETION_GATE", "0")

    bid_path = tmp_path / "bid.docx"
    doc = Document()
    doc.add_paragraph("商务投标文件")
    doc.add_paragraph("资格审查申请书")
    doc.add_paragraph("申请人：测试科技有限公司")
    doc.save(str(bid_path))

    coarse_output = json.dumps(
        {
            "requirements": [{"id": "R001", "category": "响应格式", "text": "t", "source": "s"}],
            "findings": [
                {
                    "id": "F001",
                    "requirement_id": "R001",
                    "status": "non_compliant",
                    "issue": "申请人填写错误",
                    "tender_evidence": "招标文件第1页 L1-L2：资格审查申请书应由投标人填写",
                    "bid_evidence": "投标文件主体信息错误",
                    "recommendation": "更正申请人名称",
                }
            ],
            "summary": {"requirement_count": 1, "finding_count": 1},
        },
        ensure_ascii=False,
    )
    precise_output = json.dumps(
        {
            "requirements": [{"id": "R001", "category": "响应格式", "text": "t", "source": "s"}],
            "findings": [
                {
                    "id": "F001",
                    "requirement_id": "R001",
                    "status": "non_compliant",
                    "issue": "申请人填写错误",
                    "tender_evidence": "招标文件第1页 L1-L2：资格审查申请书应由投标人填写",
                    "bid_evidence": "投标文件《资格审查申请书》L2-L2：申请人：测试科技有限公司",
                    "recommendation": "更正申请人名称",
                }
            ],
            "summary": {"requirement_count": 1, "finding_count": 1},
        },
        ensure_ascii=False,
    )

    client = _SeqClient(outputs=[coarse_output, precise_output])

    report, raw = run_bid_review_with_claude(
        tender_path=str(tmp_path / "tender.pdf"),
        bid_path=str(bid_path),
        client=client,
        extra_instruction="",
        user_instruction="",
    )

    assert client.ask_text_calls == 2
    assert "[证据精确定位强制要求]" in client.prompts[1]
    assert report["findings"][0]["bid_evidence"].startswith("投标文件《资格审查申请书》L2-L2：")
    assert "[LOCATION_RETRY]" in raw


def test_location_retry_merges_precise_evidence_without_dropping_findings(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv("BID_REVIEW_ENABLE_SECOND_PASS", "0")
    monkeypatch.setenv("BID_REVIEW_DOCX_OCR_REQUIRED", "0")
    monkeypatch.setenv("BID_REVIEW_ENFORCE_COMPLETION_GATE", "0")

    bid_path = tmp_path / "bid.docx"
    doc = Document()
    doc.add_paragraph("无关内容")
    doc.save(str(bid_path))

    coarse_output = json.dumps(
        {
            "requirements": [{"id": "R001", "category": "响应格式", "text": "t", "source": "s"}],
            "findings": [
                {
                    "id": "F001",
                    "requirement_id": "R001",
                    "status": "non_compliant",
                    "issue": "申请人填写错误",
                    "tender_evidence": "招标文件第1页 L1-L2：资格审查申请书应由投标人填写",
                    "bid_evidence": "投标文件主体信息错误A",
                    "recommendation": "更正申请人名称",
                },
                {
                    "id": "F002",
                    "requirement_id": "R001",
                    "status": "risk",
                    "issue": "授权代表信息缺失",
                    "tender_evidence": "招标文件第2页 L1-L2：授权代表信息应完整",
                    "bid_evidence": "投标文件主体信息错误B",
                    "recommendation": "补充授权代表信息",
                },
            ],
            "summary": {"requirement_count": 1, "finding_count": 2},
        },
        ensure_ascii=False,
    )
    retry_output = json.dumps(
        {
            "requirements": [{"id": "R001", "category": "响应格式", "text": "t", "source": "s"}],
            "findings": [
                {
                    "id": "F001",
                    "requirement_id": "R001",
                    "status": "non_compliant",
                    "issue": "申请人填写错误",
                    "tender_evidence": "招标文件第1页 L1-L2：资格审查申请书应由投标人填写",
                    "bid_evidence": "投标文件《资格审查申请书》L2-L2：申请人：测试科技有限公司",
                    "recommendation": "更正申请人名称",
                }
            ],
            "summary": {"requirement_count": 1, "finding_count": 1},
        },
        ensure_ascii=False,
    )

    client = _SeqClient(outputs=[coarse_output, retry_output])

    report, raw = run_bid_review_with_claude(
        tender_path=str(tmp_path / "tender.pdf"),
        bid_path=str(bid_path),
        client=client,
        extra_instruction="",
        user_instruction="",
    )

    assert client.ask_text_calls == 2
    assert len(report["findings"]) == 2
    assert any(item["issue"] == "授权代表信息缺失" for item in report["findings"])
    first = next(item for item in report["findings"] if item["issue"] == "申请人填写错误")
    assert first["bid_evidence"].startswith("投标文件《资格审查申请书》L2-L2：")
    assert "[LOCATION_RETRY]" in raw

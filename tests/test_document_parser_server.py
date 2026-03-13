from __future__ import annotations

import asyncio
import json
import zipfile
from pathlib import Path
from docx import Document

from app.mcp_servers import document_parser_server as server
from app.review.execution_policy import OCRFilterPolicy
from app.review.prepared_artifacts import extract_or_load_word_image_manifest


def test_extract_or_load_word_image_manifest_dedups_duplicate_images(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("BID_REVIEW_RUNTIME_ROOT", str(tmp_path / "runtime"))
    docx_path = tmp_path / "sample.docx"
    with zipfile.ZipFile(docx_path, "w") as zf:
        zf.writestr("[Content_Types].xml", "")
        zf.writestr("word/media/image1.png", b"duplicate-image-content")
        zf.writestr("word/media/image2.png", b"duplicate-image-content")

    artifact = extract_or_load_word_image_manifest(
        docx_path,
        filter_policy=OCRFilterPolicy(min_image_bytes=1, min_image_edge_px=0, dedup_enabled=True),
    ).data

    assert artifact["image_count_raw"] == 2
    assert artifact["image_count_unique"] == 1
    assert len(artifact["image_paths"]) == 1
    assert sum(1 for item in artifact["images"] if item["is_duplicate"]) == 1


def test_get_pdf_outline_returns_cached_outline(monkeypatch, tmp_path: Path) -> None:
    fake_data = {
        "file_hash": "abc",
        "cache_hit": True,
        "outline": {
            "total_pages": 5,
            "sections": [{"title": "第二章 投标人须知", "page_no": 2}],
            "relevant_sections": ["第二章 投标人须知"],
        },
    }

    class _FakeBuild:
        def __init__(self) -> None:
            self.data = fake_data

    monkeypatch.setattr(server, "load_or_build_tender_artifact", lambda _path: _FakeBuild())

    result = json.loads(asyncio.run(server.get_pdf_outline("D:/fake.pdf")))

    assert result["file_hash"] == "abc"
    assert result["cache_hit"] is True
    assert result["total_pages"] == 5


def test_search_word_text_returns_section_hits(monkeypatch) -> None:
    fake_data = {
        "sections": [
            {
                "id": "1",
                "title": "投标函",
                "lines": [
                    {"line_no": 1, "text": "投标函", "norm": "投标函"},
                    {"line_no": 2, "text": "致：内蒙古昆明卷烟有限责任公司", "norm": "致内蒙古昆明卷烟有限责任公司"},
                ],
            }
        ]
    }

    class _FakeBuild:
        def __init__(self) -> None:
            self.data = fake_data

    monkeypatch.setattr(server, "load_or_build_bid_artifact", lambda _path: _FakeBuild())

    result = json.loads(asyncio.run(server.search_word_text("D:/fake.docx", "昆明卷烟", 20, 1)))

    assert result["hits"][0]["title"] == "投标函"
    assert result["hits"][0]["hits"][0]["line_no"] == 2


def test_search_word_text_supports_pipe_or_query(monkeypatch) -> None:
    fake_data = {
        "sections": [
            {
                "id": "1",
                "title": "投标函",
                "lines": [
                    {"line_no": 1, "text": "投标函", "norm": "投标函"},
                    {"line_no": 2, "text": "致：内蒙古昆明卷烟有限责任公司", "norm": "致内蒙古昆明卷烟有限责任公司"},
                ],
            }
        ]
    }

    class _FakeBuild:
        def __init__(self) -> None:
            self.data = fake_data

    monkeypatch.setattr(server, "load_or_build_bid_artifact", lambda _path: _FakeBuild())

    result = json.loads(asyncio.run(server.search_word_text("D:/fake.docx", "开标一览表|昆明卷烟", 20, 1)))

    assert result["hits"][0]["title"] == "投标函"


def test_get_word_outline_skips_toc_like_sections(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("BID_REVIEW_RUNTIME_ROOT", str(tmp_path / "runtime"))
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("1、 投标函 9")
    doc.add_paragraph("2、 开标一览表 12")
    doc.add_paragraph("投标函")
    doc.add_paragraph("致：内蒙古昆明卷烟有限责任公司")
    doc.add_paragraph("开标一览表")
    doc.add_paragraph("项目名称：测试项目")
    doc.save(docx_path)

    result = json.loads(asyncio.run(server.get_word_outline(str(docx_path))))
    titles = [item["title"] for item in result["sections"]]

    assert "1、 投标函 9" not in titles
    assert "2、 开标一览表 12" not in titles
    assert "投标函" in titles
    assert "开标一览表" in titles


def test_safe_word_image_output_dir_forces_runtime_temp_root(monkeypatch, tmp_path: Path) -> None:
    runtime_root = tmp_path / "runtime"
    monkeypatch.setenv("BID_REVIEW_RUNTIME_ROOT", str(runtime_root))

    safe_dir = server._safe_word_image_output_dir("D:/code/docs/extracted_images2")

    assert safe_dir is not None
    assert str(safe_dir).startswith(str((runtime_root / "tmp" / "document-parser").resolve()))
    assert str(safe_dir).endswith("extracted_images2")
